"""Service layer for Product Prices Directory."""

from __future__ import annotations

import io
import uuid
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BadRequestException
from app.core.logging import get_logger
from app.masters.import_export import build_csv_export, build_excel_export
from app.masters.product_prices.repository import ProductPriceRepository
from app.masters.product_prices.schemas import (
    AssignSupplierPricePayload,
    PriceImportSummary,
    ProductPriceItem,
    ProductPriceSupplierItem,
    UpdatePricePayload,
)

logger = get_logger(__name__)


class ProductPriceService:
    """Service managing product prices and supplier links."""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.repo = ProductPriceRepository(session)

    async def list_prices(
        self,
        *,
        page: int = 1,
        page_size: int = 50,
        search: str | None = None,
        category_id: uuid.UUID | None = None,
        sub_category_id: uuid.UUID | None = None,
        brand_id: uuid.UUID | None = None,
        has_price: bool | None = None,
        sort_by: str = "product_name_tally",
        sort_dir: str = "asc",
    ) -> tuple[list[ProductPriceItem], int]:
        return await self.repo.list_product_prices(
            page=page,
            page_size=page_size,
            search=search,
            category_id=category_id,
            sub_category_id=sub_category_id,
            brand_id=brand_id,
            has_price=has_price,
            sort_by=sort_by,
            sort_dir=sort_dir,
        )

    async def get_product_suppliers(self, product_id: uuid.UUID) -> list[ProductPriceSupplierItem]:
        return await self.repo.get_product_suppliers(product_id)

    async def list_suppliers_lookup(self) -> list[dict[str, Any]]:
        return await self.repo.list_suppliers_lookup()

    async def assign_price(self, payload: AssignSupplierPricePayload) -> uuid.UUID:
        link_id = await self.repo.assign_supplier_price(payload)
        await self.session.flush()
        return link_id

    async def update_price(self, link_id: uuid.UUID, payload: UpdatePricePayload) -> None:
        await self.repo.update_supplier_price(link_id, payload)
        await self.session.flush()

    async def delete_price(self, link_id: uuid.UUID) -> None:
        await self.repo.delete_supplier_price(link_id)
        await self.session.flush()

    async def export_prices(
        self,
        file_format: str = "xlsx",
        search: str | None = None,
        category_id: uuid.UUID | None = None,
        sub_category_id: uuid.UUID | None = None,
        brand_id: uuid.UUID | None = None,
        has_price: bool | None = None,
    ) -> tuple[bytes, str, str]:
        """Export products with their pricing details (supports filtering and professional openpyxl styling)."""
        items, _ = await self.repo.list_product_prices(
            page=1,
            page_size=50000,
            search=search,
            category_id=category_id,
            sub_category_id=sub_category_id,
            brand_id=brand_id,
            has_price=has_price,
            sort_by="product_name_tally",
            sort_dir="asc",
        )
        headers = [
            "Sr. No.",
            "Product Code",
            "Product Name (Tally)",
            "Category",
            "Sub-Category",
            "Brand",
            "UOM",
            "Best Price",
            "Currency",
            "Primary Supplier",
            "Total Quoting Suppliers",
            "Pricing Status",
        ]

        if file_format.lower() == "csv":
            rows: list[dict[str, Any]] = []
            for idx, item in enumerate(items, start=1):
                has_quote = item.best_price is not None
                rows.append(
                    {
                        "Sr. No.": idx,
                        "Product Code": item.product_code or "",
                        "Product Name (Tally)": item.product_name_tally or item.product_name,
                        "Category": item.category_name or "",
                        "Sub-Category": item.sub_category_name or "",
                        "Brand": item.brand_name or "",
                        "UOM": item.uom_code or "",
                        "Best Price": f"{item.best_price:.2f}" if has_quote else "Unpriced",
                        "Currency": item.best_currency or "CNY",
                        "Primary Supplier": item.primary_supplier_name or "",
                        "Total Quoting Suppliers": item.supplier_count,
                        "Pricing Status": "Priced" if has_quote else "Unpriced",
                    }
                )
            csv_bytes = build_csv_export(headers, rows)
            return csv_bytes, "text/csv", "product_prices.csv"

        # Styled Excel Workbook
        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet(title="Product Prices")
        else:
            ws.title = "Product Prices"
        ws.views.sheetView[0].showGridLines = True

        ws.append(headers)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Corporate Dark Navy
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        thin_side = Side(border_style="thin", color="CBD5E1")
        border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)

        ws.row_dimensions[1].height = 26

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border

        even_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        data_font = Font(name="Calibri", size=10, color="1E293B")

        for row_idx, item in enumerate(items, start=2):
            ws.row_dimensions[row_idx].height = 20
            row_fill = even_fill if row_idx % 2 == 0 else odd_fill
            has_quote = item.best_price is not None
            best_price_val: float | str = float(item.best_price) if item.best_price is not None else "Unpriced"

            row_data = [
                row_idx - 1,
                item.product_code or "—",
                item.product_name_tally or item.product_name or "—",
                item.category_name or "—",
                item.sub_category_name or "—",
                item.brand_name or "—",
                item.uom_code or "—",
                best_price_val,
                item.best_currency or "CNY",
                item.primary_supplier_name or "—",
                item.supplier_count,
                "Priced" if has_quote else "Unpriced",
            ]
            ws.append(row_data)

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.fill = row_fill
                c.font = data_font
                c.border = border

                if col_idx in (1, 2, 7, 9, 11, 12):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 8:
                    if has_quote:
                        c.alignment = Alignment(horizontal="right", vertical="center")
                        c.number_format = "#,##0.00"
                    else:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        c.font = Font(name="Calibri", size=10, italic=True, color="94A3B8")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for col in ws.columns:
            max_len = 0
            col_idx = col[0].column
            if not isinstance(col_idx, int):
                continue
            col_letter = get_column_letter(col_idx)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ws.column_dimensions["A"].width = 9   # Sr. No.
        ws.column_dimensions["B"].width = 18  # Product Code
        ws.column_dimensions["C"].width = 38  # Product Name
        ws.column_dimensions["H"].width = 15  # Best Price
        ws.column_dimensions["J"].width = 30  # Primary Supplier

        buffer = io.BytesIO()
        wb.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "product_prices.xlsx",
        )

    async def generate_template(self) -> tuple[bytes, str, str]:
        """Generate sample import template."""
        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet(title="Price Import Template")
        else:
            ws.title = "Price Import Template"

        headers = [
            "Product Code*",
            "Product Name",
            "Supplier Name*",
            "Unit Price*",
            "Currency",
            "MOQ",
            "Notes",
        ]
        ws.append(headers)

        # Sample rows
        ws.append(["DAR-01758", "100 Gm Horizontal Cutter", "Yiwu Machining Works", 125.50, "CNY", 10, "Direct factory price"])
        ws.append(["DAR-01758", "100 Gm Horizontal Cutter", "Zhejiang Precision Tools", 120.00, "CNY", 50, "Bulk discount MOQ 50"])

        buffer = io.BytesIO()
        wb.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "product_price_import_template.xlsx",
        )

    async def import_prices(self, file_bytes: bytes) -> PriceImportSummary:
        """Universal bulk Excel import for product prices."""
        try:
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            if ws is None:
                raise BadRequestException("The uploaded spreadsheet has no active sheet.")
        except BadRequestException:
            raise
        except Exception as exc:
            raise BadRequestException(f"Failed to read Excel file: {exc}") from exc

        rows_iter = ws.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            raise BadRequestException("The uploaded spreadsheet is empty.") from None

        if not raw_headers:
            raise BadRequestException("The uploaded spreadsheet has no headers.")

        # Clean headers
        col_map: dict[str, int] = {}
        for idx, h in enumerate(raw_headers):
            if h is None:
                continue
            key = str(h).strip().lower().replace("*", "").replace("_", " ").replace("-", " ")
            if "code" in key and "product" in key:
                col_map["product_code"] = idx
            elif "product" in key or "item" in key:
                col_map["product_name"] = idx
            elif "supplier" in key or "vendor" in key:
                col_map["supplier_name"] = idx
            elif "price" in key or "rate" in key or "unit" in key:
                col_map["price"] = idx
            elif "curr" in key:
                col_map["currency"] = idx
            elif "moq" in key:
                col_map["moq"] = idx
            elif "note" in key or "remark" in key:
                col_map["notes"] = idx

        if "price" not in col_map:
            raise BadRequestException("Missing required column 'Price' or 'Unit Price'.")
        if "supplier_name" not in col_map:
            raise BadRequestException("Missing required column 'Supplier Name'.")
        if "product_code" not in col_map and "product_name" not in col_map:
            raise BadRequestException("Missing required column 'Product Code' or 'Product Name'.")

        # Cache products and suppliers for ultra-fast in-memory resolution
        prod_res = await self.session.execute(
            text("SELECT id, product_code, product_name, product_name_tally FROM products WHERE deleted_at IS NULL;")
        )
        prod_by_code: dict[str, uuid.UUID] = {}
        prod_by_name: dict[str, uuid.UUID] = {}
        for r in prod_res.fetchall():
            pid = r[0]
            if r[1]:
                prod_by_code[str(r[1]).strip().lower()] = pid
            if r[2]:
                prod_by_name[str(r[2]).strip().lower()] = pid
            if r[3]:
                prod_by_name[str(r[3]).strip().lower()] = pid

        supp_res = await self.session.execute(
            text("SELECT id, company_name FROM suppliers WHERE deleted_at IS NULL;")
        )
        supp_by_name: dict[str, uuid.UUID] = {}
        for r in supp_res.fetchall():
            sid = r[0]
            if r[1]:
                supp_by_name[str(r[1]).strip().lower()] = sid

        summary = PriceImportSummary()
        row_num = 1  # 1 is header

        for row in rows_iter:
            row_num += 1
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            summary.total_rows += 1

            # 1. Resolve Product
            product_id: uuid.UUID | None = None
            prod_code_raw = (
                str(row[col_map["product_code"]]).strip() if "product_code" in col_map and row[col_map["product_code"]] else ""
            )
            prod_name_raw = (
                str(row[col_map["product_name"]]).strip() if "product_name" in col_map and row[col_map["product_name"]] else ""
            )

            if prod_code_raw:
                product_id = prod_by_code.get(prod_code_raw.lower())
            if not product_id and prod_name_raw:
                product_id = prod_by_name.get(prod_name_raw.lower())

            if not product_id:
                summary.failed += 1
                summary.errors.append({
                    "row": row_num,
                    "error": f"Product not found: Code '{prod_code_raw}', Name '{prod_name_raw}'",
                })
                continue

            # 2. Resolve Supplier
            supplier_id: uuid.UUID | None = None
            supp_name_raw = (
                str(row[col_map["supplier_name"]]).strip() if row[col_map["supplier_name"]] else ""
            )
            if supp_name_raw:
                supplier_id = supp_by_name.get(supp_name_raw.lower())

            if not supplier_id:
                summary.failed += 1
                summary.errors.append({
                    "row": row_num,
                    "error": f"Supplier not found: '{supp_name_raw}'",
                })
                continue

            # 3. Parse Unit Price
            price_raw = row[col_map["price"]]
            try:
                price_val = float(str(price_raw).replace(",", "").replace("¥", "").replace("$", "").strip())
                if price_val < 0:
                    raise ValueError("Price must be non-negative")
            except Exception:
                summary.failed += 1
                summary.errors.append({
                    "row": row_num,
                    "error": f"Invalid unit price: '{price_raw}'",
                })
                continue

            # 4. Optional fields
            curr_val = "CNY"
            if "currency" in col_map and row[col_map["currency"]]:
                curr_val = str(row[col_map["currency"]]).strip().upper()[:10]

            moq_val = None
            if "moq" in col_map and row[col_map["moq"]] is not None:
                try:
                    moq_raw = str(row[col_map["moq"]]).strip()
                    if moq_raw:
                        moq_val = float(moq_raw)
                except Exception:
                    pass

            notes_val = None
            if "notes" in col_map and row[col_map["notes"]]:
                notes_val = str(row[col_map["notes"]]).strip()

            payload = AssignSupplierPricePayload(
                product_id=product_id,
                supplier_id=supplier_id,
                unit_price=price_val,
                currency=curr_val,
                moq=moq_val,
                notes=notes_val,
            )

            try:
                await self.repo.assign_supplier_price(payload)
                summary.created += 1
            except Exception as exc:
                summary.failed += 1
                summary.errors.append({"row": row_num, "error": str(exc)})

        await self.session.flush()
        return summary
